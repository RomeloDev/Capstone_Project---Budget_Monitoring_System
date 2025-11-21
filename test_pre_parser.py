"""
Test script to analyze PRE Excel template and detect custom line items
"""

import openpyxl
from decimal import Decimal

def analyze_pre_template(file_path):
    """Analyze the PRE Excel template structure"""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb.active

    print("=" * 80)
    print("PRE EXCEL TEMPLATE ANALYSIS")
    print("=" * 80)

    # 1. Find all rows with data in columns E-I (Q1-Total)
    print("\n1. SCANNING ALL DATA ROWS (with Q1-Q4 values):")
    print("-" * 80)

    data_rows = []
    for row_num in range(1, ws.max_row + 1):
        item_name = ws[f'A{row_num}'].value
        q1 = ws[f'E{row_num}'].value
        q2 = ws[f'F{row_num}'].value
        q3 = ws[f'G{row_num}'].value
        q4 = ws[f'H{row_num}'].value
        total = ws[f'I{row_num}'].value

        # Check if row has any quarterly data
        has_data = any([
            q1 and str(q1).strip().upper() not in ['', 'NONE', '-'],
            q2 and str(q2).strip().upper() not in ['', 'NONE', '-'],
            q3 and str(q3).strip().upper() not in ['', 'NONE', '-'],
            q4 and str(q4).strip().upper() not in ['', 'NONE', '-'],
        ])

        if has_data and item_name:
            data_rows.append({
                'row': row_num,
                'item_name': item_name,
                'q1': q1,
                'q2': q2,
                'q3': q3,
                'q4': q4,
                'total': total
            })
            print(f"Row {row_num:3d}: {item_name[:50]:<50} | Q1={q1}")

    print(f"\nTotal data rows found: {len(data_rows)}")

    # 2. Find the Grand Total row
    print("\n2. LOCATING GRAND TOTAL ROW:")
    print("-" * 80)

    grand_total_row = None
    for row_num in range(1, ws.max_row + 1):
        cell_value = ws[f'A{row_num}'].value
        if cell_value and str(cell_value).strip().upper() == 'TOTAL':
            # Check if this is the grand total (not a subtotal)
            # Grand total should be after all categories
            if row_num > 170:  # Assuming grand total is near the end
                grand_total_row = row_num
                print(f"Grand Total found at Row {row_num}")
                print(f"  Q1: {ws[f'E{row_num}'].value}")
                print(f"  Q2: {ws[f'F{row_num}'].value}")
                print(f"  Q3: {ws[f'G{row_num}'].value}")
                print(f"  Q4: {ws[f'H{row_num}'].value}")
                print(f"  Total: {ws[f'I{row_num}'].value}")
                break

    # 3. Identify section boundaries
    print("\n3. SECTION BOUNDARIES:")
    print("-" * 80)

    sections = {}
    section_markers = [
        ('RECEIPTS', 9, 10),
        ('PERSONNEL', 13, 18),
        ('MOOE_START', 20, 21),
        ('MOOE_END', 130, 131),
        ('CAPITAL_START', 133, 134),
        ('CAPITAL_END', 174, 175),
    ]

    for marker_name, start, end in section_markers:
        print(f"{marker_name}: Rows {start}-{end}")
        if start <= ws.max_row:
            print(f"  Row {start}: {ws[f'A{start}'].value}")

    # 4. Detect potential custom items (rows not in parser mapping)
    print("\n4. DETECTING POTENTIAL CUSTOM ITEMS:")
    print("-" * 80)
    print("These are rows with data that might not be in the parser's CELL_MAPPINGS")

    # Import the parser to check which rows are mapped
    import sys
    sys.path.append('apps/end_user_app/utils')
    from pre_parser import PREParser

    # Get all mapped cell positions
    mapped_rows = set()

    # Extract row numbers from receipts
    for item in PREParser.CELL_MAPPINGS['receipts']:
        row = int(item[1][1:])  # e.g., 'E10' -> 10
        mapped_rows.add(row)

    # Extract row numbers from personnel
    for item in PREParser.CELL_MAPPINGS['personnel']:
        row = int(item[1][1:])
        mapped_rows.add(row)

    # Extract row numbers from MOOE
    for category, items in PREParser.CELL_MAPPINGS['mooe'].items():
        for item in items:
            row = int(item[1][1:])
            mapped_rows.add(row)

    # Extract row numbers from capital
    for category, items in PREParser.CELL_MAPPINGS['capital'].items():
        for item in items:
            row = int(item[1][1:])
            mapped_rows.add(row)

    print(f"Total mapped rows in parser: {len(mapped_rows)}")
    print(f"Mapped rows: {sorted(mapped_rows)}")

    # Find unmapped data rows
    unmapped_rows = []
    for data_row in data_rows:
        if data_row['row'] not in mapped_rows:
            unmapped_rows.append(data_row)

    print(f"\nPotential unmapped/custom rows: {len(unmapped_rows)}")
    for urow in unmapped_rows:
        print(f"  Row {urow['row']}: {urow['item_name'][:60]}")

    # 5. Check for blank rows where custom items could be inserted
    print("\n5. BLANK ROWS (potential custom item insertion points):")
    print("-" * 80)

    blank_ranges = []
    in_blank = False
    blank_start = None

    for row_num in range(10, 177):  # Between first item and grand total
        is_blank = not ws[f'A{row_num}'].value and not ws[f'E{row_num}'].value

        if is_blank and not in_blank:
            blank_start = row_num
            in_blank = True
        elif not is_blank and in_blank:
            blank_ranges.append((blank_start, row_num - 1))
            in_blank = False

    print(f"Found {len(blank_ranges)} blank row ranges:")
    for start, end in blank_ranges:
        if end - start >= 2:  # Only show ranges of 3+ rows
            print(f"  Rows {start}-{end} ({end-start+1} rows)")

    return {
        'data_rows': data_rows,
        'grand_total_row': grand_total_row,
        'mapped_rows': mapped_rows,
        'unmapped_rows': unmapped_rows,
        'blank_ranges': blank_ranges
    }


if __name__ == '__main__':
    result = analyze_pre_template('excel_templates/Departmental-PRE.xlsx')

    print("\n" + "=" * 80)
    print("ANALYSIS COMPLETE")
    print("=" * 80)
    print(f"Total data rows: {len(result['data_rows'])}")
    print(f"Mapped in parser: {len(result['mapped_rows'])}")
    print(f"Unmapped rows: {len(result['unmapped_rows'])}")
    print(f"Grand total row: {result['grand_total_row']}")

    print("\n" + "=" * 80)
    print("RECOMMENDATION FOR CUSTOM ITEMS")
    print("=" * 80)
    print("""
1. DYNAMIC ROW DETECTION APPROACH:
   Instead of fixed cell mappings, scan all rows between section headers
   and grand total to capture ALL items (including custom ones).

2. SECTION DETECTION:
   - Detect section headers (Personnel, MOOE, Capital)
   - Scan all rows within each section
   - Identify items by checking if columns E-I have data
   - Stop at subtotal rows

3. CUSTOM ITEM IDENTIFICATION:
   - Any row with data that's not in the standard template
   - User can insert rows anywhere in the appropriate section
   - Parser should dynamically detect and extract all rows

4. VALIDATION:
   - Check that grand total matches sum of all items
   - Verify quarterly totals match (Q1+Q2+Q3+Q4 = Total)
   - Ensure custom items are in correct sections
    """)
