"""
Script to create BISU Header Template for Reports
This creates a reusable Excel template with BISU header (Row 1) only
"""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.drawing.image import Image
import os

# Create new workbook
wb = Workbook()
ws = wb.active
ws.title = "Report"

# BISU header text (5 lines)
header_text = """Republic of the Philippines
BOHOL ISLAND STATE UNIVERSITY
Magsija, Balilihan, 6342, Bohol, Philippines
Office of the Administration and Finance
Balance I Integrity I Stewardship I Uprightness"""

# 1. Merge cells A1:I1
ws.merge_cells('A1:I1')

# 2. Set cell value and formatting
cell = ws['A1']
cell.value = header_text
cell.font = Font(name='Arial', size=11, bold=False)
cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# 3. Add medium bottom border
border = Border(bottom=Side(style='medium', color='FF000000'))
for col in range(1, 10):
    ws.cell(1, col).border = border

# 4. Set row height
ws.row_dimensions[1].height = 84.75

# 5. Set column widths
column_widths = {
    'A': 13.14, 'B': 8.43, 'C': 10.57, 'D': 10.71,
    'E': 10.71, 'F': 10.71, 'G': 10.71, 'H': 13.71, 'I': 10.71
}
for col_letter, width in column_widths.items():
    ws.column_dimensions[col_letter].width = width

# 6. Add logos
logo_dir = os.path.join('apps', 'end_user_app', 'static', 'logos')

# Check if logos exist
if os.path.exists(logo_dir):
    # Left logo (BISU seal)
    seal_path = os.path.join(logo_dir, 'bisu_seal.png')
    if os.path.exists(seal_path):
        img_seal = Image(seal_path)
        img_seal.width = 90
        img_seal.height = 90
        ws.add_image(img_seal, 'A1')
        print(f"[OK] Added BISU seal logo")
    else:
        print(f"[WARNING] BISU seal not found at {seal_path}")

    # Right logos
    bagong_path = os.path.join(logo_dir, 'bagong_pilipinas.png')
    if os.path.exists(bagong_path):
        img_bagong = Image(bagong_path)
        img_bagong.width = 90
        img_bagong.height = 90
        ws.add_image(img_bagong, 'I1')
        print(f"[OK] Added Bagong Pilipinas logo")
    else:
        print(f"[WARNING] Bagong Pilipinas logo not found at {bagong_path}")

    iso_path = os.path.join(logo_dir, 'iso_cert.png')
    if os.path.exists(iso_path):
        img_iso = Image(iso_path)
        img_iso.width = 80
        img_iso.height = 80
        ws.add_image(img_iso, 'H1')
        print(f"[OK] Added ISO certification logo")
    else:
        print(f"[WARNING] ISO cert logo not found at {iso_path}")
else:
    print(f"[ERROR] Logo directory not found: {logo_dir}")
    print("The template will be created without logos")

# Save template
template_dir = os.path.join('apps', 'end_user_app', 'templates', 'excel_templates')
os.makedirs(template_dir, exist_ok=True)

template_path = os.path.join(template_dir, 'BISU_Report_Template.xlsx')
wb.save(template_path)

print(f"\n[SUCCESS] BISU Report Template created at:")
print(f"  {template_path}")
print(f"\nThis template contains:")
print(f"  - Row 1: BISU Header with logos and official text")
print(f"  - Proper formatting (fonts, alignment, borders)")
print(f"  - Column widths matching Departmental-PRE.xlsx")
print(f"\nUse this template for all budget reports in the system!")
