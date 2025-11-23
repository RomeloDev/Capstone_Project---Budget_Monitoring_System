import openpyxl
from openpyxl.utils import get_column_letter

# Load the template
wb = openpyxl.load_workbook('../excel_templates/Departmental-PRE.xlsx')
ws = wb.active

print("="*70)
print("ANALYZING Departmental-PRE.xlsx - ROW 1 (Columns A to I)")
print("="*70)

# Check merged cells in Row 1
print("\n=== MERGED CELLS IN ROW 1 ===")
for merged in ws.merged_cells.ranges:
    if merged.min_row == 1:
        print(f"  {merged}")

# Check each column A-I
print("\n=== EACH CELL IN ROW 1 ===")
for col in range(1, 10):  # A to I (1 to 9)
    col_letter = get_column_letter(col)
    cell = ws[f'{col_letter}1']

    print(f"\n{col_letter}1:")
    print(f"  Value: {repr(cell.value)[:100] if cell.value else 'None'}")
    print(f"  Font: {cell.font.name}, Size: {cell.font.size}, Bold: {cell.font.bold}")
    print(f"  Alignment: H={cell.alignment.horizontal}, V={cell.alignment.vertical}, Wrap={cell.alignment.wrap_text}")

    # Check for fill color
    if hasattr(cell.fill, 'fgColor') and cell.fill.fgColor:
        print(f"  Fill: {cell.fill.fgColor.rgb}")
    else:
        print(f"  Fill: None")

    # Border
    if cell.border.bottom and cell.border.bottom.style:
        print(f"  Bottom Border: {cell.border.bottom.style}")

# Check for images/logos
print("\n=== IMAGES/LOGOS IN WORKSHEET ===")
if hasattr(ws, '_images'):
    print(f"Number of images: {len(ws._images)}")
    for idx, img in enumerate(ws._images):
        print(f"\nImage {idx+1}:")
        print(f"  Anchor: {img.anchor}")
        if hasattr(img, 'width') and hasattr(img, 'height'):
            print(f"  Size: {img.width} x {img.height}")
else:
    print("No images found or images not accessible")

# Row height
print(f"\n=== ROW 1 HEIGHT ===")
print(f"Height: {ws.row_dimensions[1].height}")

# Column widths
print(f"\n=== COLUMN WIDTHS (A-I) ===")
for col in range(1, 10):
    col_letter = get_column_letter(col)
    width = ws.column_dimensions[col_letter].width
    print(f"{col_letter}: {width}")

print("\n" + "="*70)
