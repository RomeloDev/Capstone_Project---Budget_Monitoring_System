"""
Script to create BISU Header Template from existing Departmental-PRE.xlsx
This copies Row 1 (BISU header with logos) and creates a clean template
"""
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import os

# Source file
source_file = r"C:\Users\John Romel Lucot\OneDrive\Desktop\Capstone project\bb_budget_monitoring_system\excel_templates\Departmental-PRE.xlsx"

print(f"Loading source file: {source_file}")

# Load the source workbook
wb_source = load_workbook(source_file)
ws_source = wb_source.active

# Create new workbook for template
wb_template = Workbook()
ws_template = wb_template.active
ws_template.title = "Report"

print("Copying Row 1 (BISU Header)...")

# Copy Row 1 data and formatting
for col in range(1, 10):  # A to I
    col_letter = get_column_letter(col)

    # Copy cell value
    source_cell = ws_source[f'{col_letter}1']
    target_cell = ws_template[f'{col_letter}1']

    if source_cell.value:
        target_cell.value = source_cell.value

    # Copy formatting
    if source_cell.font:
        target_cell.font = source_cell.font.copy()
    if source_cell.alignment:
        target_cell.alignment = source_cell.alignment.copy()
    if source_cell.border:
        target_cell.border = source_cell.border.copy()
    if source_cell.fill:
        target_cell.fill = source_cell.fill.copy()

# Copy merged cells for Row 1
for merged_range in ws_source.merged_cells.ranges:
    if merged_range.min_row == 1 and merged_range.max_row == 1:
        ws_template.merge_cells(str(merged_range))
        print(f"  Merged cells: {merged_range}")

# Copy row height
ws_template.row_dimensions[1].height = ws_source.row_dimensions[1].height
print(f"  Row height: {ws_source.row_dimensions[1].height}")

# Copy column widths
for col in range(1, 10):
    col_letter = get_column_letter(col)
    ws_template.column_dimensions[col_letter].width = ws_source.column_dimensions[col_letter].width

print("  Column widths copied")

# Copy images (logos)
print("Copying logos...")
if hasattr(ws_source, '_images') and ws_source._images:
    from copy import copy as make_copy
    for image in ws_source._images:
        try:
            # Just copy the image object directly
            ws_template.add_image(make_copy(image))
            print(f"  Copied image")
        except Exception as e:
            print(f"  Warning: Could not copy image: {e}")

# Save template in templates/excel_templates directory
template_dir = os.path.join('apps', 'end_user_app', 'templates', 'excel_templates')
os.makedirs(template_dir, exist_ok=True)

template_path = os.path.join(template_dir, 'BISU_Report_Template.xlsx')
wb_template.save(template_path)

print(f"\n[SUCCESS] BISU Report Template created at:")
print(f"  {template_path}")
print(f"\nTemplate contains:")
print(f"  - Row 1: Complete BISU Header with logos")
print(f"  - Original formatting and styling")
print(f"  - Ready to use for all budget reports!")
print(f"\nThis template will be used by both:")
print(f"  - end_user_app reports")
print(f"  - admin_panel reports")
