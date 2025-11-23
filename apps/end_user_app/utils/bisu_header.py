"""
BISU Header Utility for Budget Reports

This module provides functions to add the official BISU header
(with logos and institutional text) to Excel reports.

Template-based approach: Uses BISU_Report_Template.xlsx for consistency
"""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import os
from django.conf import settings
from copy import copy as make_copy


def add_bisu_header_to_excel(ws, start_row=1):
    """
    Add official BISU header to Excel worksheet using template.
    This ensures perfect consistency with the original Departmental-PRE.xlsx format.

    The template (BISU_Report_Template.xlsx) contains:
    - Row 1: Complete BISU header with logos and text
    - Merged cells A1:I1
    - BISU seal logo (left side, column A)
    - Bagong Pilipinas and ISO logos (right side, columns H-I)
    - Official 5-line institutional text (centered)
    - Proper formatting, borders, and column widths

    Args:
        ws: openpyxl Worksheet object
        start_row: Row number to insert header (default: 1)

    Returns:
        int: Next available row after header

    Raises:
        FileNotFoundError: If BISU_Report_Template.xlsx not found
    """
    # Load BISU template
    template_path = os.path.join(
        settings.BASE_DIR,
        'apps', 'end_user_app', 'templates', 'excel_templates',
        'BISU_Report_Template.xlsx'
    )

    if not os.path.exists(template_path):
        raise FileNotFoundError(
            f"BISU Report Template not found at {template_path}. "
            "Run create_bisu_template_from_existing.py to create it."
        )

    # Load template
    wb_template = load_workbook(template_path)
    ws_template = wb_template.active

    # Copy Row 1 from template to target worksheet
    # 1. Copy cell values and formatting (columns A to I)
    for col in range(1, 10):
        col_letter = get_column_letter(col)

        source_cell = ws_template[f'{col_letter}1']
        target_cell = ws[f'{col_letter}{start_row}']

        # Copy value
        if source_cell.value:
            target_cell.value = source_cell.value

        # Copy formatting
        if source_cell.font:
            target_cell.font = make_copy(source_cell.font)
        if source_cell.alignment:
            target_cell.alignment = make_copy(source_cell.alignment)
        if source_cell.border:
            target_cell.border = make_copy(source_cell.border)
        if source_cell.fill:
            target_cell.fill = make_copy(source_cell.fill)

    # 2. Copy merged cells
    for merged_range in ws_template.merged_cells.ranges:
        if merged_range.min_row == 1 and merged_range.max_row == 1:
            # Adjust to target row
            if start_row == 1:
                ws.merge_cells(str(merged_range))
            else:
                # Adjust range for different start row
                ws.merge_cells(
                    start_row=start_row,
                    start_column=merged_range.min_col,
                    end_row=start_row,
                    end_column=merged_range.max_col
                )

    # 3. Copy row height
    ws.row_dimensions[start_row].height = ws_template.row_dimensions[1].height

    # 4. Copy column widths
    for col in range(1, 10):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = ws_template.column_dimensions[col_letter].width

    # 5. Copy images (logos)
    if hasattr(ws_template, '_images') and ws_template._images:
        for image in ws_template._images:
            # Create a copy of the image
            img_copy = make_copy(image)

            # Add to target worksheet
            # Note: Images will be positioned at their original anchor points
            ws.add_image(img_copy)

    return start_row + 1


def get_bisu_header_context():
    """
    Get context data for rendering BISU header in HTML templates.

    Used for HTML preview of reports before exporting.

    Returns:
        dict: Context with header text and logo paths for HTML rendering
    """
    return {
        'bisu_header': {
            'line1': 'Republic of the Philippines',
            'line2': 'BOHOL ISLAND STATE UNIVERSITY',
            'line3': 'Magsija, Balilihan, 6342, Bohol, Philippines',
            'line4': 'Office of the Administration and Finance',
            'line5': 'Balance I Integrity I Stewardship I Uprightness',
        },
        'logo_bisu_seal': 'logos/bisu_seal.png',
        'logo_bagong_pilipinas': 'logos/bagong_pilipinas.png',
        'logo_iso_cert': 'logos/iso_cert.png',
    }
